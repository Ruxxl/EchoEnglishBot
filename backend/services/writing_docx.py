"""Отчёт Writing Checker в .docx: баллы, отзыв, эссе с правками прямо в тексте и образец ответа."""

import io

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt, RGBColor

from backend.services.writing import TASK_TYPES

_CRITERIA_TITLES = {
    "task_response": "Task Response / Achievement",
    "coherence_cohesion": "Coherence & Cohesion",
    "lexical_resource": "Lexical Resource",
    "grammar_accuracy": "Grammatical Range & Accuracy",
}
_RED = RGBColor(0xC6, 0x28, 0x28)
_GREEN = RGBColor(0x1B, 0x7F, 0x3B)
_GREY = RGBColor(0x6B, 0x7A, 0x8A)


def _bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def build_report(check: dict) -> bytes:
    result = check["result"]
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    doc.add_heading("IELTS Writing — отчёт о проверке", level=0)
    meta = doc.add_paragraph()
    meta.add_run(f"{TASK_TYPES[check['task_type']].split(' (')[0]} · {result['word_count']} слов · ")
    meta.add_run(check["created_at"].strftime("%d.%m.%Y")).italic = True
    if check["topic"]:
        p = doc.add_paragraph()
        p.add_run("Задание: ").bold = True
        p.add_run(check["topic"])

    p = doc.add_paragraph()
    run = p.add_run(f"Overall Band: {result['overall_band']:.1f}")
    run.bold, run.font.size = True, Pt(20)

    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    for cell, text in zip(table.rows[0].cells, ("Критерий", "Балл", "Комментарий")):
        cell.text = text
    for key, title in _CRITERIA_TITLES.items():
        crit = result["criteria"][key]
        row = table.add_row().cells
        row[0].text, row[1].text = title, f"{crit['band']:.1f}"
        row[2].text = crit["comment"] + (f"\n\n{crit['reasoning']}" if crit.get("reasoning") else "")

    if result["summary"]:
        doc.add_heading("Общий отзыв", level=1)
        doc.add_paragraph(result["summary"])
    if result["strengths"]:
        doc.add_heading("Сильные стороны", level=2)
        _bullets(doc, result["strengths"])
    if result["improvements"]:
        doc.add_heading("Что улучшить", level=2)
        _bullets(doc, result["improvements"])
    if result.get("teacher_comment"):
        doc.add_heading("Комментарий преподавателя", level=1)
        doc.add_paragraph(result["teacher_comment"])

    # Эссе с правками в тексте: зачёркнутое красным — как было, зелёным — как надо.
    doc.add_heading("Эссе с исправлениями", level=1)
    essay, issues = check["essay"], result["issues"]
    for paragraph_text, offset in _paragraphs_with_offsets(essay):
        p = doc.add_paragraph()
        cursor, end = offset, offset + len(paragraph_text)
        for issue in issues:
            if issue["start"] < offset or issue["end"] > end:
                continue
            p.add_run(essay[cursor : issue["start"]])
            wrong = p.add_run(issue["original"])
            wrong.font.strike, wrong.font.color.rgb = True, _RED
            fix = p.add_run(f" {issue['suggestion']}")
            fix.bold, fix.font.color.rgb = True, _GREEN
            fix.font.highlight_color = WD_COLOR_INDEX.BRIGHT_GREEN if issue["category"] == "word_choice" else None
            cursor = issue["end"]
        p.add_run(essay[cursor:end])

    if issues:
        doc.add_heading("Разбор ошибок", level=1)
        for n, issue in enumerate(issues, 1):
            p = doc.add_paragraph()
            p.add_run(f"{n}. ").bold = True
            wrong = p.add_run(issue["original"])
            wrong.font.color.rgb = _RED
            p.add_run(" → ")
            fix = p.add_run(issue["suggestion"])
            fix.bold, fix.font.color.rgb = True, _GREEN
            tag = p.add_run(f"  [{issue['category']}]")
            tag.font.size, tag.font.color.rgb = Pt(9), _GREY
            if issue["explanation"]:
                p.add_run(f"\n{issue['explanation']}").italic = True

    if result.get("sample_essay"):
        doc.add_heading("Образец ответа (Band 8+)", level=1)
        for text, _ in _paragraphs_with_offsets(result["sample_essay"]):
            doc.add_paragraph(text)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _paragraphs_with_offsets(text: str) -> list[tuple[str, int]]:
    out, offset = [], 0
    for line in text.split("\n"):
        if line.strip():
            out.append((line, offset))
        offset += len(line) + 1
    return out
